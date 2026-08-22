from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("outputs") / "SOC已有研究基础_开题报告适配版.docx"


def set_font(run, font_name="宋体"):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def add_run(paragraph, text, *, bold=False, size=11, color=None):
    run = paragraph.add_run(text)
    set_font(run)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def add_para(doc, text="", *, first_line=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(6)
    if first_line:
        p.paragraph_format.first_line_indent = Pt(22)
    if text:
        add_run(p, text)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(6)
    color = (31, 78, 121) if level == 1 else (64, 64, 64)
    size = 15 if level == 1 else 13
    add_run(p, text, bold=True, size=size, color=color)
    return p


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def build():
    OUT.parent.mkdir(exist_ok=True)
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(title, "SOC 已有研究基础补充稿", bold=True, size=18, color=(31, 78, 121))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(subtitle, "适配《锂离子电池 SOC-SOH 联合估计方法研究》开题报告", size=11, color=(89, 89, 89))

    add_heading(doc, "一、总体定位", 1)
    add_para(
        doc,
        "SOC 部分建议在开题报告中定位为已经形成阶段性成果的机理模型支路。"
        "已有工作以二阶 RC 等效电路模型为基础，结合 FFRLS 在线参数辨识和 GST-AEKF 门控强跟踪自适应扩展卡尔曼滤波算法，"
        "实现动态工况下的 SOC 在线估计和端电压预测。该成果可作为后续 SOC-SOH 联合估计中的 SOC 快变量估计核心。"
    )
    add_para(
        doc,
        "写作时不建议把重点放在多种滤波器谁最优的数值比较上，而应突出 GST-AEKF 的算法机制："
        "激进 Q 自适应更新用于提高误差响应速度，NIS 门控用于识别并抑制异常测量，强跟踪因子用于增强工况突变后的状态跟踪能力。"
        "其他滤波算法可作为对比方法出现，但不应削弱 GST-AEKF 作为已有投稿成果和 SOC 支路基准算法的地位。"
    )

    add_heading(doc, "二、可补充到 4.1 的内容", 1)
    add_para(doc, "建议放在 4.1“基于机理模型的 SOC 估计模块”末尾，用于说明 SOC 模块已有基础。", first_line=False)
    add_para(
        doc,
        "在已有 SOC 研究基础方面，已完成基于二阶 RC 等效电路模型和 GST-AEKF 的 SOC 估计方法研究。"
        "该方法以 SOC 和两个 RC 极化支路电压作为状态变量，通过 OCV-SOC 多项式关系和端电压观测方程建立非线性状态空间模型，"
        "并结合 FFRLS 在线参数辨识方法动态更新欧姆内阻和极化参数。"
        "在滤波算法方面，GST-AEKF 在 AEKF 基础上引入激进 Q 自适应更新、NIS 门控机制和强跟踪因子，"
        "使滤波器能够在动态工况和异常测量条件下实现更快的误差响应和更稳定的状态修正。"
        "该 SOC 估计框架已完成动态工况下的仿真验证，并已形成阶段性投稿成果，可作为本课题 SOC-SOH 联合估计中的机理模型支路基础。"
    )

    add_heading(doc, "三、可补充到 4.4 的内容", 1)
    add_para(doc, "建议替换或扩充 4.4 中“第一，SOC 单模块验证”这一段。", first_line=False)
    add_para(
        doc,
        "第一，SOC 单模块验证。已有工作已基于公开锂离子电池动态工况数据完成 SOC 模块实验验证，"
        "覆盖不同温度条件、不同动态工况和不同初始 SOC 场景。实验重点验证二阶 RC 等效电路模型、"
        "FFRLS 在线参数辨识和 GST-AEKF 滤波估计框架的有效性，评价指标包括 SOC 估计 RMSE、MAE、最大误差、"
        "端电压预测误差以及初始 SOC 偏差收敛能力。结果表明，GST-AEKF 通过激进 Q 自适应更新、"
        "NIS 门控和强跟踪因子的协同作用，能够提升动态工况下 SOC 估计的响应速度和鲁棒性，"
        "为后续引入 SOH 健康参数反馈提供了可靠的 SOC 在线估计基础。"
    )

    add_heading(doc, "四、可补充到第 6 节的内容", 1)
    add_para(doc, "建议作为第 6 节“已有研究进度”中的 SOC 段，与 SOH 段和联合框架段并列。", first_line=False)
    add_para(
        doc,
        "在 SOC 估计方面，已完成基于 2RC 等效电路模型、FFRLS 在线参数辨识和 GST-AEKF 的 SOC 估计框架设计与仿真验证。"
        "已有工作建立了 OCV-SOC 多项式关系、2RC 状态空间模型和端电压观测方程，并在 AEKF 基础上引入激进 Q 自适应更新、"
        "NIS 门控和强跟踪因子，形成门控强跟踪自适应扩展卡尔曼滤波算法。"
        "该方法能够在动态工况和异常测量条件下增强滤波器的响应速度和鲁棒性，相关阶段性研究成果已整理投稿。"
        "上述工作为后续 SOC-SOH 联合估计中的 SOC 机理模型支路和容量反馈修正机制奠定了基础。"
    )

    add_heading(doc, "五、与 SOC-SOH 联合估计的衔接", 1)
    add_para(
        doc,
        "GST-AEKF 与 SOC-SOH 联合估计框架的衔接重点在于健康参数反馈。"
        "首先，SOH 支路输出的当前最大可用容量可替换 SOC 状态方程中的固定额定容量，修正电池老化导致的安时积分偏差。"
        "其次，SOH 或健康因子可用于修正 2RC 模型中的欧姆内阻和极化参数，使端电压观测模型更符合老化阶段的实际响应。"
        "最后，现有 OCV-SOC 多项式观测方程可进一步扩展为随 SOH 变化的 OCV-SOC 修正模型。"
        "因此，GST-AEKF 可作为 SOC 在线估计器，接收 SOH 支路提供的低频健康状态反馈，形成“SOH 慢变量更新 + SOC 快变量递推”的双时间尺度联合估计框架。"
    )

    add_heading(doc, "六、建议插入位置速览", 1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, text in enumerate(["位置", "插入内容", "作用"]):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        add_run(p, text, bold=True, size=10.5)
        shade_cell(hdr[i], "D9EAF7")

    rows = [
        ("4.1 末尾", "GST-AEKF 算法基础", "说明 SOC 支路已有投稿成果"),
        ("4.4 第一层实验", "SOC 单模块验证", "说明已完成动态工况 SOC 验证"),
        ("6 已有研究进度", "SOC 已有基础段", "与 SOH 和联合框架并列"),
        ("联合框架说明", "容量、内阻、OCV-SOC 反馈接口", "衔接 SOC-SOH 联合估计主线"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            add_run(p, text, size=10.5)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(footer, "SOC 已有研究基础补充稿", size=9, color=(128, 128, 128))

    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(OUT.resolve())
