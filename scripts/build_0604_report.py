from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\Projects\1111-soh")
TEMPLATE = Path(r"C:\Users\46139\Desktop\论文\汇报\0521汇报_v6_no_ci.docx")
OUT = ROOT / "outputs" / "0604汇报.docx"


def clear_body(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def set_run(run, size=11, bold=False, color="000000"):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc, text="", size=11, bold=False, color="000000", before=0, after=4, line=1.15):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, color=color)
    return p


def add_section(doc, title):
    return add_para(doc, title, size=12, bold=True, color="000000", before=8, after=5)


def set_cell_text(cell, text, bold=False, fill=None):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_run(r, size=10, bold=bold)
    if fill:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)


def set_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "B7C0CC")


def set_table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
            tc_pr = row.cells[idx]._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def add_result_table(doc):
    headers = ["场景", "严重程度", "模型", "AUC", "Det@FPR5%", "Delay"]
    rows = [
        ["容量拐点", "轻度", "CNN-LSTM", "0.720", "62.6%", "96.0"],
        ["容量拐点", "轻度", "PI-MS-CNNLSTM", "0.837", "77.9%", "23.0"],
        ["容量拐点", "中度", "CNN-LSTM", "0.741", "67.2%", "59.0"],
        ["容量拐点", "中度", "PI-MS-CNNLSTM", "0.869", "85.4%", "15.2"],
        ["容量拐点", "重度", "CNN-LSTM", "0.750", "68.8%", "67.8"],
        ["容量拐点", "重度", "PI-MS-CNNLSTM", "0.858", "83.9%", "21.2"],
    ]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_borders(table)
    set_table_widths(table, [1.0, 0.8, 1.55, 0.7, 1.0, 0.75])
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, fill="E8EEF5")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    return table


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(TEMPLATE)
    clear_body(doc)

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    add_para(doc, "PI-CNNLSTM 电池 SOH 估计  ·  6.04汇报", size=16, bold=True, after=8)
    add_para(
        doc,
        "本周汇报聚焦异常退化风险提示：从 HUST 77 块电池容量退化轨迹中筛选衰减速度偏快的电池作为异常样本，"
        "以 trajectory_deviation（SOH 衰减趋势）作为核心检测指标，验证已训练 SOH 估计模型是否能够在不额外训练故障诊断模型的前提下，对快衰减电池形成可解释的异常提示。",
        after=8,
    )

    add_section(doc, "一、异常电池筛选口径")
    add_para(
        doc,
        "HUST 数据集包含 77 组锂离子电池完整生命周期容量记录，不同电池在寿命长度、容量衰减速率和局部波动幅度上存在明显差异。"
        "本周按照容量退化轨迹的下降速度进行初步筛选，将衰减速度明显快于整体平均水平的一批电池视为异常退化电池。"
        "这里的“异常”并不等同于已知故障标签，而是从 SOH 轨迹形态出发定义的快衰减样本，用于检验模型对异常退化趋势的敏感性。",
    )
    add_para(
        doc,
        "● 汇报口径：正常电池表现为 SOH 随循环次数缓慢、连续下降；异常电池表现为衰减斜率增大、容量拐点提前或局部阶段性偏离正常退化趋势。",
        after=2,
    )
    add_para(
        doc,
        "● 工程意义：该筛选方式贴近 BMS 中“先发现异常退化趋势，再决定是否复核容量标定”的使用逻辑，适合作为低成本风险提示入口。",
        after=6,
    )

    add_section(doc, "二、检测指标：trajectory_deviation ★SOH衰减趋势")
    add_para(
        doc,
        "本周采用 battery_anomaly_report.docx 中整理的 trajectory_deviation 指标作为主要检测信号。该指标利用最近 20 个窗口的 SOH 预测值进行一阶线性外推，"
        "再计算当前预测值相对于外推趋势的偏离程度。正常退化时，SOH 预测轨迹通常平滑且近似单调；当电池进入快衰减或异常退化阶段时，实际预测值会偏离原有局部趋势，异常分数随之升高。",
    )
    add_para(
        doc,
        "● 指标优势：trajectory_deviation 不依赖真实 SOH 标签，只依赖模型自身预测序列，因此可在实际部署中逐循环实时计算。",
        after=2,
    )
    add_para(
        doc,
        "● 检测重点：该指标不仅关注单点骤降，也关注一段时间内衰减趋势是否变陡，因而更适合识别容量拐点、突发加速老化和析锂台阶突降等轨迹扰动明显的场景。",
        after=6,
    )

    add_section(doc, "三、异常检测结果与现象")
    add_para(
        doc,
        "根据第四章_0601修改.docx 中的实验结果，PI-MS-CNNLSTM 在异常退化场景下能够输出更平滑的正常退化本底轨迹，使异常段的 trajectory_deviation 分数更加突出。"
        "在突发加速老化场景中，异常分数在故障注入后快速超过阈值并持续维持；在容量拐点和析锂场景中，异常分数在注入点附近出现明显升高，说明 SOH 预测轨迹已经偏离原有衰减趋势。",
    )
    add_result_table(doc)
    add_para(
        doc,
        "表1  容量拐点/快衰减场景下的 trajectory_deviation 检测结果。Det@FPR5% 表示在 5% 误报率约束下的检出率，Delay 表示首次满足报警条件的窗口延迟。",
        size=10,
        after=6,
    )
    add_para(
        doc,
        "从表1可以看出，PI-MS-CNNLSTM 在轻度、中度、重度容量拐点场景下的 AUC 和 Det@FPR5% 均高于 CNN-LSTM，且报警延迟显著缩短。"
        "其中中度容量拐点场景下，PI-MS-CNNLSTM 的 Det@FPR5% 达到 85.4%，较 CNN-LSTM 的 67.2% 提高 18.2 个百分点，Delay 由 59.0 个窗口缩短至 15.2 个窗口。"
        "这说明物理一致性约束降低了正常轨迹的本底噪声，从而放大了快衰减阶段的轨迹偏差信号。",
    )

    add_section(doc, "四、下周工作计划")
    add_para(
        doc,
        "● 进一步明确 HUST 快衰减异常电池的筛选阈值，补充每块异常电池的衰减斜率、寿命终止循环和 trajectory_deviation 峰值统计。",
        after=2,
    )
    add_para(
        doc,
        "● 将异常检测结果整理为论文第四章工程扩展部分的补充材料，统一 AUC、Det@FPR5%、Delay 三项指标的表述口径。",
        after=2,
    )
    add_para(
        doc,
        "● 针对渐进型异常检测较弱的问题，尝试延长趋势外推窗口或融合 rate_anomaly、drop_anomaly 信号，提高对缓慢偏离型退化的鲁棒性。",
        after=6,
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
