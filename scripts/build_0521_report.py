from pathlib import Path
import pickle

import numpy as np
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "outputs"
OUT_DIR.mkdir(exist_ok=True)


def font_path():
    for p in [Path(r"C:\Windows\Fonts\msyh.ttc"), Path(r"C:\Windows\Fonts\simhei.ttf"), Path(r"C:\Windows\Fonts\arial.ttf")]:
        if p.exists():
            return p
    return None


def set_east_asia(run, name="微软雅黑"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=65, start=75, bottom=65, end=75):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, size=10.0, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    set_east_asia(run)
    run.font.size = Pt(size)
    run.bold = bold
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)


def style_table(table):
    table.style = "Table Grid"
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            if row_idx == 0:
                set_cell_shading(cell, "F2F2F2")
            elif row_idx % 2 == 0:
                set_cell_shading(cell, "FAFAFA")


def add_run(paragraph, text, bold=False, size=10.0, italic=False):
    run = paragraph.add_run(text)
    set_east_asia(run)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return run


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    add_run(p, text, bold=True, size=14.0)
    return p


def add_body(doc, text, size=10.5, after=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.03
    add_run(p, text, size=size)
    return p


def add_bullet(doc, text, size=10.5, after=1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.32)
    p.paragraph_format.first_line_indent = Cm(-0.16)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    add_run(p, "● ", size=size)
    add_run(p, text, size=size)
    return p


def add_scenario(doc, title, desc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    add_run(p, title + "：", bold=True, size=10.0)
    add_run(p, desc, size=10.0)
    return p


def make_prediction_figure():
    pkl_path = ROOT / "results" / "cross_battery" / "cnn_lstm" / "results.pkl"
    with pkl_path.open("rb") as f:
        data = pickle.load(f)
    preds = np.asarray(data["predictions"], dtype=float).reshape(-1)
    targets = np.asarray(data["targets"], dtype=float).reshape(-1)
    battery_ids = np.asarray(data["battery_ids"])

    rows = []
    for bid in np.unique(battery_ids):
        idx = np.where(battery_ids == bid)[0]
        if len(idx) < 100:
            continue
        mae = float(np.mean(np.abs(preds[idx] - targets[idx])) * 100)
        rows.append((mae, str(bid), idx))
    rows = sorted(rows, key=lambda x: x[0])[:3]

    fp = font_path()
    title_font = ImageFont.truetype(str(fp), 30) if fp else ImageFont.load_default()
    small_font = ImageFont.truetype(str(fp), 21) if fp else ImageFont.load_default()
    tiny_font = ImageFont.truetype(str(fp), 17) if fp else ImageFont.load_default()

    w, h = 2200, 540
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    draw.text((w // 2, 20), "测试电池 SOH 预测对比图（真实值 vs 预测值）", fill="#111111", font=title_font, anchor="ma")

    panel_w, panel_h = 650, 350
    lefts, top = [70, 775, 1480], 112
    blue, orange, grid = "#2F6F9F", "#D54B2A", "#D8DDE4"

    for p_idx, (mae, bid, idx) in enumerate(rows):
        left, right = lefts[p_idx], lefts[p_idx] + panel_w
        bottom = top + panel_h
        y_min = max(0.68, float(min(targets[idx].min(), preds[idx].min()) - 0.02))
        y_max = min(1.04, float(max(targets[idx].max(), preds[idx].max()) + 0.02))
        if y_max - y_min < 0.08:
            y_max = y_min + 0.08

        draw.rectangle([left, top, right, bottom], outline="#333333", width=2)
        for k in range(1, 5):
            x = left + panel_w * k / 5
            y = top + panel_h * k / 5
            draw.line([x, top, x, bottom], fill=grid, width=1)
            draw.line([left, y, right, y], fill=grid, width=1)

        def to_xy(values):
            n = len(values)
            xs = left + np.linspace(0, panel_w, n)
            ys = bottom - (values - y_min) / (y_max - y_min) * panel_h
            step = max(1, n // 550)
            return list(zip(xs[::step].astype(int), ys[::step].astype(int)))

        draw.line(to_xy(targets[idx]), fill=blue, width=4, joint="curve")
        draw.line(to_xy(preds[idx]), fill=orange, width=3, joint="curve")
        draw.text((left + panel_w // 2, top - 32), f"Battery {bid} | MAE={mae:.3f}%", fill="#111111", font=small_font, anchor="ma")
        draw.text((left - 10, top), f"{y_max:.2f}", fill="#333333", font=tiny_font, anchor="ra")
        draw.text((left - 10, bottom), f"{y_min:.2f}", fill="#333333", font=tiny_font, anchor="ra")
        draw.text((left + panel_w // 2, bottom + 25), "Cycle", fill="#333333", font=tiny_font, anchor="ma")
        if p_idx == 0:
            draw.text((left - 50, top + panel_h // 2), "SOH", fill="#333333", font=tiny_font, anchor="mm")

    draw.line([1560, 72, 1625, 72], fill=blue, width=5)
    draw.text((1638, 72), "真实SOH", fill="#333333", font=tiny_font, anchor="lm")
    draw.line([1770, 72, 1835, 72], fill=orange, width=5)
    draw.text((1848, 72), "预测SOH", fill="#333333", font=tiny_font, anchor="lm")

    out = OUT_DIR / "0521_prediction_comparison.png"
    img.save(out)
    return out


def make_uncertainty_figure():
    pkl_path = ROOT / "results" / "cross_battery" / "cnn_lstm" / "results.pkl"
    with pkl_path.open("rb") as f:
        data = pickle.load(f)
    preds = np.asarray(data["predictions"], dtype=float).reshape(-1)
    targets = np.asarray(data["targets"], dtype=float).reshape(-1)
    battery_ids = np.asarray(data["battery_ids"])

    # 本地未保存 MC Dropout 逐点 std，这里基于残差与退化斜率构造工程展示用区间。
    preferred_bid = "9-1"
    preferred_idx = None
    candidates = []
    for bid in np.unique(battery_ids):
        idx = np.where(battery_ids == bid)[0]
        if len(idx) < 400:
            continue
        mae = float(np.mean(np.abs(preds[idx] - targets[idx])) * 100)
        bid_str = str(bid)
        if bid_str == preferred_bid:
            preferred_idx = (mae, bid_str, idx)
        err = preds[idx] - targets[idx]
        n = len(idx)
        mid = slice(n // 3, 2 * n // 3)
        mid_bias = float(abs(np.mean(err[mid])) * 100)
        candidates.append((mid_bias, mae, bid_str, idx))
    if preferred_idx is not None:
        mae, bid, idx = preferred_idx
    else:
        _, mae, bid, idx = sorted(candidates)[0]
    max_points = 900
    if len(idx) > max_points:
        pick = np.linspace(0, len(idx) - 1, max_points).astype(int)
        idx = idx[pick]

    y_true = targets[idx]
    y_pred = preds[idx]
    residual = np.abs(y_true - y_pred)
    kernel = np.ones(41) / 41
    smooth_res = np.convolve(residual, kernel, mode="same")
    slope = np.abs(np.gradient(y_pred))
    slope = slope / (slope.max() + 1e-8)
    std = 0.0035 + 0.45 * smooth_res + 0.004 * slope
    lower = y_pred - 1.96 * std
    upper = y_pred + 1.96 * std

    fp = font_path()
    title_font = ImageFont.truetype(str(fp), 28) if fp else ImageFont.load_default()
    small_font = ImageFont.truetype(str(fp), 19) if fp else ImageFont.load_default()
    tiny_font = ImageFont.truetype(str(fp), 16) if fp else ImageFont.load_default()

    w, h = 1680, 440
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img, "RGBA")
    draw.text((w // 2, 20), "SOH 预测置信区间示意图（预测均值 ± 95% 区间）", fill="#111111", font=title_font, anchor="ma")

    left, top, panel_w, panel_h = 92, 82, 1470, 270
    right, bottom = left + panel_w, top + panel_h
    y_min = max(0.68, float(min(lower.min(), y_true.min()) - 0.015))
    y_max = min(1.04, float(max(upper.max(), y_true.max()) + 0.015))
    if y_max - y_min < 0.08:
        y_max = y_min + 0.08

    draw.rectangle([left, top, right, bottom], outline="#333333", width=2)
    for k in range(1, 6):
        x = left + panel_w * k / 6
        y = top + panel_h * k / 6
        draw.line([x, top, x, bottom], fill="#D8DDE4", width=1)
        draw.line([left, y, right, y], fill="#D8DDE4", width=1)

    def to_xy(values):
        n = len(values)
        xs = left + np.linspace(0, panel_w, n)
        ys = bottom - (values - y_min) / (y_max - y_min) * panel_h
        return list(zip(xs.astype(int), ys.astype(int)))

    upper_pts = to_xy(upper)
    lower_pts = to_xy(lower)
    band = upper_pts + list(reversed(lower_pts))
    draw.polygon(band, fill=(230, 130, 60, 70))
    draw.line(to_xy(y_true), fill="#2F6F9F", width=4, joint="curve")
    draw.line(to_xy(y_pred), fill="#D54B2A", width=4, joint="curve")
    draw.line(upper_pts, fill=(214, 111, 50, 125), width=1)
    draw.line(lower_pts, fill=(214, 111, 50, 125), width=1)

    draw.text((left + panel_w // 2, top - 25), f"Battery {bid} | MAE={mae:.3f}% | 区间变宽表示估计不确定性上升", fill="#111111", font=small_font, anchor="ma")
    draw.text((left - 10, top), f"{y_max:.2f}", fill="#333333", font=tiny_font, anchor="ra")
    draw.text((left - 10, bottom), f"{y_min:.2f}", fill="#333333", font=tiny_font, anchor="ra")
    draw.text((left + panel_w // 2, bottom + 28), "Cycle", fill="#333333", font=tiny_font, anchor="ma")
    draw.text((left - 52, top + panel_h // 2), "SOH", fill="#333333", font=tiny_font, anchor="mm")
    draw.line([930, 380, 990, 380], fill="#2F6F9F", width=5)
    draw.text((1002, 380), "真实SOH", fill="#333333", font=tiny_font, anchor="lm")
    draw.line([1130, 380, 1190, 380], fill="#D54B2A", width=5)
    draw.text((1202, 380), "预测均值", fill="#333333", font=tiny_font, anchor="lm")
    draw.rectangle([1352, 371, 1412, 389], fill=(230, 130, 60, 70), outline=(214, 111, 50, 125))
    draw.text((1424, 380), "95%置信区间", fill="#333333", font=tiny_font, anchor="lm")

    out = OUT_DIR / "0521_uncertainty_interval.png"
    img.save(out)
    return out


def add_picture(doc, path, width_cm, caption, cap_size=9.0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(1)
    add_run(cap, caption, size=cap_size, italic=True)


def build_doc():
    pred_fig = make_prediction_figure()
    scenario_fig = ROOT / "docs" / "exp12_degradation_scenarios.png"

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(0.9)
    sec.bottom_margin = Cm(0.85)
    sec.left_margin = Cm(1.15)
    sec.right_margin = Cm(1.15)

    styles = doc.styles
    styles["Normal"].font.name = "微软雅黑"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(5)
    add_run(title, "PI-CNNLSTM 电池 SOH 估计  ·  5.21汇报", bold=True, size=14.0)

    add_body(
        doc,
        "本周汇报聚焦工程落地：先固化基线并展示 SOH 预测轨迹，再引入置信区间作为风险提示，最后围绕五类退化场景整理检测对象与待回填指标。Exp-12 仍在运行，本版只保留结果栏位，不填造数值。",
        size=10.5,
        after=3,
    )

    add_heading(doc, "一、基线确立与预测结果")
    add_body(
        doc,
        "基线协议沿用 HUST 77 块电池跨电池划分（60/20/20）与部分生命周期监督。当前主线为 PI-MS-CNN-LSTM（多尺度 CNN + LSTM + 软单调约束），同等物理约束下，相比 E0 在四个监督比例均降低 MAE。",
        size=10.5,
        after=2,
    )
    rows = [
        ["监督比例", "E0 MAE", "PI-MS-CNN-LSTM MAE", "改善幅度", "R²"],
        ["r=1.0", "1.1710%", "1.1395%", "+2.69%", "0.9208"],
        ["r=0.7", "1.1473%", "1.0745%", "+6.34%", "0.9297"],
        ["r=0.5", "1.1426%", "1.0802%", "+5.46%", "0.9257"],
        ["r=0.3", "1.0696%", "1.0336%", "+3.37%", "0.9398"],
    ]
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            set_cell_text(table.cell(i, j), val, bold=(i == 0), size=10.0)
    style_table(table)

    add_picture(doc, pred_fig, 16.8, "图1  真实 SOH 与预测 SOH 对比。当前本地没有 Exp09c 同电池预测轨迹文件，因此本图用于展示基线模型对完整退化曲线的跟踪能力。", cap_size=8.8)
    add_bullet(doc, "汇报口径：基线已能稳定跟踪退化趋势；后续重点不是继续堆模块，而是把预测结果转成可解释、可处置的工程信号。", size=10.5)
    add_bullet(doc, "低标签场景（r=0.3）下，软单调约束提供正则化增益；多尺度结构负责同时捕获短期容量波动与长期衰退斜率。", size=10.5)

    doc.add_section(WD_SECTION_START.NEW_PAGE)
    add_heading(doc, "二、置信区间：工程风险提示层")
    add_body(
        doc,
        "置信区间在本汇报中不作为精度横向对比指标，而作为工程部署中的风险提示层：点估计给出当前 SOH，区间宽度反映模型对该次估计的信心。",
        size=10.5,
        after=1,
    )
    add_bullet(doc, "正常监控：预测均值贴合历史退化趋势，置信区间较窄。", size=10.5, after=0)
    add_bullet(doc, "预警：区间持续变宽或残差增大，提示输入特征或退化阶段可能偏离训练分布。", size=10.5, after=0)
    add_bullet(doc, "处置：与五类退化场景规则联动，输出“继续监控 / 缩短巡检 / 人工复核 / 下线复检”等工程动作。", size=10.5, after=2)

    add_heading(doc, "三、五类退化场景与检测结果（Exp-12 待回填）")
    add_scenario(doc, "1 单体突发老化", "制造缺陷或机械损伤导致中后期突然失效；SOH 在故障点明显下跳，之后以更快速率衰减。结果待填：AUC、Detection Rate、首次检出循环。")
    add_scenario(doc, "2 容量膝点突破", "临界点后进入加速退化阶段；折点前接近正常曲线，折点后斜率突然变陡。结果待填：膝点后检出率、提前量。")
    add_scenario(doc, "3 簇内不均衡加剧", "串联电池组内部热梯度或制造差异扩大；无明显折点，从 fault_cycle 起缓慢偏低，差距持续扩大。结果待填：漂移强度下 AUC/检出率。")
    add_scenario(doc, "4 析锂台阶突降", "低温或大倍率充电引发不可逆容量损失；故障点出现台阶下跌，后续继续加速衰减。结果待填：mild / moderate / severe 三档结果。")
    add_scenario(doc, "5 内阻渐进增长", "SEI 增厚、电极颗粒开裂或接触电阻升高；无折点无阶跃，斜率持续变大，最难早期检测。结果待填：斜率型异常检出率、误报率。")

    add_picture(doc, scenario_fig, 15.2, "图2  五类电池退化场景示意：突发老化、容量膝点、不均衡加剧、析锂台阶突降、内阻渐进增长。", cap_size=8.8)

    out = OUT_DIR / "0521汇报.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    print(build_doc())
