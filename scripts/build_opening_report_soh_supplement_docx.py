from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("outputs") / "开题报告_SOH已有研究补充建议.docx"


def set_east_asian_font(run, font_name="宋体"):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def add_run(paragraph, text, *, bold=False, size=11, color=None):
    run = paragraph.add_run(text)
    set_east_asian_font(run)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def add_para(doc, text="", *, first_line=True):
    p = doc.add_paragraph()
    if text:
        add_run(p, text)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(6)
    if first_line:
        p.paragraph_format.first_line_indent = Pt(22)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(6)
    add_run(
        p,
        text,
        bold=True,
        size=15 if level == 1 else 13,
        color=(31, 78, 121) if level == 1 else (64, 64, 64),
    )
    return p


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, *, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(text) < 20 else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    run = add_run(p, text, bold=bold, size=10.5)
    return run


def build():
    OUT.parent.mkdir(exist_ok=True)
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "开题报告 SOH 已有研究补充建议", bold=True, size=18, color=(31, 78, 121))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "适用于《锂离子电池 SOC-SOH 联合估计方法研究》开题报告修订", size=11, color=(89, 89, 89))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "审阅范围：当前目录内 SOH 项目成果；SOC 项目材料暂不纳入。", size=10.5, color=(89, 89, 89))

    add_heading(doc, "一、总体判断")
    add_para(doc, "当前开题报告的主线是“机理模型估计 SOC + 数据驱动估计 SOH + 健康参数反馈修正 SOC”。这个框架是成立的，但其中 SOH 部分仍偏概念化，和当前项目中已经完成的 SOH 工作不完全匹配。")
    add_para(doc, "建议将 SOH 部分从“准备采用 LSTM/GRU/CNN-LSTM 开展研究”调整为“已围绕部分生命周期监督场景完成多尺度物理一致性 SOH 估计方法研究”。这样既能体现已有工作量，也能支撑 SOC-SOH 联合估计中 SOH 支路的成熟性。")

    add_heading(doc, "二、建议优先修改的位置")

    add_heading(doc, "1. 4.2 基于数据驱动的 SOH 估计模块", level=2)
    add_para(doc, "当前写法偏计划性，例如“本课题拟采用 NASA 和 CALCE 公开数据集开展 SOH 建模与验证”。建议保留 NASA/CALCE 作为后续联合估计扩展验证数据源，同时明确当前已经完成的 HUST SOH 研究基础。")
    add_para(doc, "建议补充文本如下：", first_line=False)
    add_para(doc, "在已有 SOH 研究基础方面，本课题已基于 HUST 公开锂离子电池退化数据集开展了较为系统的 SOH 估计研究。该数据集包含 77 块 LFP/石墨电池单体的循环老化数据，能够反映不同电池个体在容量衰退速率、退化阶段和寿命终点上的差异。围绕实际 BMS 中 SOH 标定不连续、全生命周期标签难以完整获得的问题，已有工作将 SOH 估计建模为部分生命周期监督任务，即保留完整运行特征轨迹，但仅在部分循环区间计算 SOH 标签监督损失。")
    add_para(doc, "在模型方面，已有研究构建了多尺度物理一致性 CNN-LSTM 模型。该模型采用多尺度并行一维卷积分支提取不同时间尺度下的充电动态特征，并结合 LSTM 捕捉跨循环退化趋势；同时引入软单调物理一致性约束，使模型输出在稳定退化阶段符合 SOH 整体不可逆衰减的物理规律。该方法既保持了深度模型对复杂非线性退化特征的表达能力，又通过物理约束提高了低标签比例场景下预测轨迹的合理性。")

    add_heading(doc, "2. 4.4 实验验证与性能评价", level=2)
    add_para(doc, "当前实验设计仍停留在 SOC/SOH/联合估计三层验证的计划阶段。建议在 SOH 单模块验证中加入已完成实验，而不是只写“拟验证”。")
    add_para(doc, "建议补充文本如下：", first_line=False)
    add_para(doc, "在 SOH 单模块验证方面，已有工作已经完成多组消融实验。实验设置覆盖完整监督及部分生命周期监督场景，监督比例包括 1.0、0.7、0.5 和 0.3，并采用多随机种子重复实验以降低偶然性影响。实验结果表明，在相同软单调物理约束下，多尺度 PI-MS-CNN-LSTM 相比单路 PI-CNN-LSTM 在各监督比例下均取得更低 MAE，提升幅度约为 2.7% 至 6.3%。其中在低监督比例 r=0.3 条件下，PI-MS-CNN-LSTM 取得 MAE=1.0336%、R²=0.9398 的结果，说明该模型能够在 SOH 标签不完整条件下保持较好的退化轨迹恢复能力。")
    add_para(doc, "同时，已有研究对注意力机制、MC Dropout、不确定性伪标签、速率连续性约束和自适应损失权重等扩展模块进行了消融分析。结果表明，复杂模块并非在低监督场景下总能带来收益，部分模块会因标签稀疏、伪标签质量不足或随机扰动引入额外退化。因此，后续联合估计研究中可优先采用结构清晰、验证稳定的 PI-MS-CNN-LSTM 作为 SOH 支路，而将不确定性量化和伪标签机制作为扩展分析或辅助模块。")

    add_heading(doc, "3. 5 预期创新点", level=2)
    add_para(doc, "当前创新点主要围绕联合估计，SOH 支路贡献没有充分体现。建议新增或改写一个创新点。")
    add_para(doc, "建议补充文本如下：", first_line=False)
    add_para(doc, "提出面向部分生命周期监督的多尺度物理一致性 SOH 估计支路。针对实际 BMS 中 SOH 标签获取成本高、完整生命周期标定困难的问题，构建 Label Masking 训练机制，在保留完整运行特征轨迹的同时仅对可用标签区间计算监督损失，并利用软单调物理一致性约束为无标签区间提供结构性弱监督。该设计能够更贴近工程中“运行数据连续可得、健康标签稀疏可得”的真实条件。")

    add_heading(doc, "4. 6 已有研究基础", level=2)
    add_para(doc, "这是最需要重写的部分。当前写法偏笼统，建议改为“SOC 基础 + SOH 已完成基础 + 联合估计接口”三段。其中 SOH 段应写得最实。")
    add_para(doc, "建议替换文本如下：", first_line=False)
    add_para(doc, "在 SOH 估计方面，已完成面向部分生命周期监督的多尺度物理一致性 SOH 估计方法研究。已有工作基于 HUST 77 块锂离子电池退化数据集，构建了循环级健康特征输入、跨电池训练测试划分、标签掩码监督和物理一致性约束训练流程。模型方面，已完成 CNN-LSTM 基线、多尺度 CNN-LSTM、物理一致性约束、多模块消融和鲁棒性验证等实验。最终确定 PI-MS-CNN-LSTM 作为 SOH 支路的主模型，其核心结构为多尺度并行 CNN 与 LSTM 组合，并引入软单调约束以增强低标签比例场景下的退化轨迹一致性。")
    add_para(doc, "实验结果表明，多尺度结构是当前 SOH 估计模型的主要性能来源。在同等软单调约束条件下，PI-MS-CNN-LSTM 相比 PI-CNN-LSTM 在不同监督比例下均取得稳定提升；在 r=0.3 的低监督场景下，软单调约束进一步提供正则化收益，使模型获得 MAE=1.0336%、R²=0.9398 的较优结果。针对注意力机制、MC Dropout、速率连续性、自适应损失权重和伪标签扩展等模块的实验也已完成，结果表明部分复杂模块在低监督条件下存在退化风险，因此后续联合估计中将以结构稳定、机制明确的 PI-MS-CNN-LSTM 作为 SOH 估计支路。")
    add_para(doc, "该 SOH 支路可直接为 SOC-SOH 联合估计提供当前可用容量估计、退化趋势信息和健康状态反馈接口。后续将在 SOC 机理模型支路基础上，优先验证 SOH 到当前最大可用容量的反馈修正，再逐步探索内阻参数和 OCV-SOC 曲线随老化状态变化的修正机制。")

    add_heading(doc, "三、需要避免的写法")
    avoid_items = [
        "不建议把 NASA/CALCE 写成“已经完成系统预处理和实验验证”。当前目录内已成体系的是 HUST SOH 项目，NASA/CALCE 更适合写成后续联合估计验证数据源。",
        "不建议把 M1-M8 全部写成最终模型组成。根据当前实验结论，最终 SOH 主线应是 PI-MS-CNN-LSTM：多尺度 CNN 是主贡献，软单调约束是低监督场景下的辅助正则化；M1/M2/M4/M5/M6 更适合写为消融边界分析。",
        "不建议把“Full Stack”作为最终方案。当前最终方案是 Exp09c，即多尺度 CNN-LSTM + 软单调约束，不包含 M2+M4。",
        "不建议把伪标签写成核心创新。M6 在 r≥0.7 有一定效果，但在 r≤0.5 不稳定，不适合作为开题报告中的主创新。",
        "文档中多个公式位置在结构化抽取时显示为空白，仅保留编号。后续改 Word 时需要重点检查公式是否真实存在，尤其是 SOC 递推公式、端电压方程、SOH 定义式、RMSE/MAE 指标公式。",
    ]
    for item in avoid_items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.first_line_indent = Pt(-18)
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(4)
        add_run(p, "• ", size=11)
        add_run(p, item, size=11)

    add_heading(doc, "四、SOH 叙事主线建议")
    add_para(doc, "建议开题报告里把 SOH 部分统一成下面这条线：")
    add_para(doc, "实际 BMS 中 SOH 标签稀疏可得，因此 SOH 估计不能默认完整生命周期监督。已有研究已将 SOH 估计建模为部分生命周期监督任务，通过 Label Masking 保留完整运行特征而只屏蔽缺失标签；在此基础上，构建 PI-MS-CNN-LSTM 模型，用多尺度 CNN 捕捉不同退化粒度特征，用 LSTM 建模跨循环趋势，用软单调约束保证低标签区间的物理一致性。实验表明，多尺度架构是主要贡献，软单调约束在极低监督比例下提供额外正则化收益。该模型可作为 SOC-SOH 联合估计中的 SOH 支路，为容量反馈、内阻修正和 OCV-SOC 修正提供健康状态输入。")

    add_heading(doc, "五、关键结果速览")
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["项目", "当前结论", "写入开题报告的建议"]
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True)
        set_cell_shading(table.rows[0].cells[i], "D9EAF7")

    rows = [
        ["数据集", "HUST 77 块电池，循环老化数据", "作为 SOH 已有研究基础重点写入"],
        ["学习场景", "部分生命周期监督，Label Masking", "作为工程真实性和方法特色写入"],
        ["主模型", "PI-MS-CNN-LSTM", "作为 SOC-SOH 联合估计中的 SOH 支路"],
        ["核心贡献", "多尺度 CNN 是主要性能来源", "不要把所有模块都写成最终模型"],
        ["物理约束", "软单调约束在 r=0.3 有正则化收益", "写成低监督场景下的辅助约束"],
        ["最佳结果", "r=0.3：MAE=1.0336%，R²=0.9398", "可作为已有实验基础数据"],
        ["不建议主写", "M6 伪标签、Full Stack、M2+M4", "放在消融或边界分析中"],
    ]
    for row_data in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row_data):
            set_cell_text(cells[i], text)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(footer, "SOH 已有研究补充建议", size=9, color=(128, 128, 128))

    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(OUT.resolve())
