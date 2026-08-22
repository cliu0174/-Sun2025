"""Extract a traceable structural/content audit from a DOCX manuscript.

The script is read-only. It prints paragraph styles, paragraph text, and table
cell text so that a LaTeX rewrite can be checked against the authoritative
Chinese draft without editing the Word file.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


def compact(text: str) -> str:
    return " ".join(text.replace("\u00a0", " ").split())


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    parser.add_argument("--headings-only", action="store_true")
    parser.add_argument("--max-chars", type=int, default=0)
    parser.add_argument("--start", type=int, default=1)
    parser.add_argument("--end", type=int, default=0)
    args = parser.parse_args()

    document = Document(args.docx)
    print(f"SOURCE\t{args.docx.resolve()}")
    print(f"PARAGRAPHS\t{len(document.paragraphs)}")
    print(f"TABLES\t{len(document.tables)}")

    for index, paragraph in enumerate(document.paragraphs, start=1):
        if index < args.start or (args.end and index > args.end):
            continue
        text = compact(paragraph.text)
        if not text:
            continue
        style = paragraph.style.name if paragraph.style is not None else ""
        if args.headings_only and not (
            "heading" in style.lower()
            or "title" in style.lower()
            or "标题" in style
        ):
            continue
        if args.max_chars and len(text) > args.max_chars:
            text = text[: args.max_chars - 1] + "…"
        print(f"P{index:04d}\t{style}\t{text}")

    if args.headings_only:
        return

    for table_index, table in enumerate(document.tables, start=1):
        print(f"TABLE\t{table_index}\tROWS={len(table.rows)}\tCOLS={len(table.columns)}")
        for row_index, row in enumerate(table.rows, start=1):
            cells = [compact(cell.text) for cell in row.cells]
            print(f"T{table_index:02d}R{row_index:03d}\t" + "\t".join(cells))


if __name__ == "__main__":
    main()
