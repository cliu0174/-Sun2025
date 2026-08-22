from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path("outputs") / "开题报告_SOH已有研究补充建议.md"
OUT = Path("outputs") / "开题报告_SOH已有研究补充建议_v3.docx"


def set_font(run, font_name="宋体"):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def add_run(paragraph, text, *, bold=False, size=11, color=None):
    run = paragraph.add_run(text)
    set_font(run)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def clean_inline(text):
    return text.replace("`", "")


def add_para(doc, text, *, quote=False, list_item=False, numbered=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(6)

    if quote:
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.first_line_indent = Pt(0)
        add_run(p, clean_inline(text), size=10.5, color=(64, 64, 64))
    elif list_item:
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.first_line_indent = Pt(-18)
        add_run(p, "• " if not numbered else "", size=11)
        add_run(p, clean_inline(text), size=11)
    else:
        p.paragraph_format.first_line_indent = Pt(22)
        add_run(p, clean_inline(text), size=11)
    return p


def add_heading(doc, text, level):
    if level == 1:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        add_run(p, clean_inline(text), bold=True, size=18, color=(31, 78, 121))
        return

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10 if level == 2 else 6)
    p.paragraph_format.space_after = Pt(6)
    size = 15 if level == 2 else 13
    color = (31, 78, 121) if level == 2 else (64, 64, 64)
    add_run(p, clean_inline(text), bold=True, size=size, color=color)


def build():
    text = SOURCE.read_text(encoding="utf-8")
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)

    pending_quote = []
    pending_para = []

    def flush_quote():
        nonlocal pending_quote
        if pending_quote:
            add_para(doc, " ".join(pending_quote).strip(), quote=True)
            pending_quote = []

    def flush_para():
        nonlocal pending_para
        if pending_para:
            add_para(doc, " ".join(pending_para).strip())
            pending_para = []

    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            flush_quote()
            flush_para()
            continue

        if line.startswith("#"):
            flush_quote()
            flush_para()
            level = len(line) - len(line.lstrip("#"))
            add_heading(doc, line[level:].strip(), level)
            continue

        if line.startswith(">"):
            flush_para()
            pending_quote.append(line.lstrip(">").strip())
            continue

        flush_quote()

        if line.startswith("- "):
            flush_para()
            add_para(doc, line[2:].strip(), list_item=True)
            continue

        if len(line) > 3 and line[0].isdigit() and line[1:3] == ". ":
            flush_para()
            add_para(doc, line.strip(), list_item=True, numbered=True)
            continue

        pending_para.append(line)

    flush_quote()
    flush_para()

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(footer, "开题报告 SOH 已有研究补充建议", size=9, color=(128, 128, 128))

    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(OUT.resolve())
