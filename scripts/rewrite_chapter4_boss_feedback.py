from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "第四章.docx"
OUT = ROOT / "第四章_老总反馈重构稿.docx"


def body_children(doc):
    return list(doc.element.body)


def child_text(child):
    return "".join(t.text or "" for t in child.iter(qn("w:t")))


def replace_in_xml(elem, replacements):
    for t in elem.iter(qn("w:t")):
        if not t.text:
            continue
        text = t.text
        for old, new in replacements.items():
            text = text.replace(old, new)
        t.text = text


def append_clone(doc, elem, replacements=None):
    new_elem = deepcopy(elem)
    if replacements:
        replace_in_xml(new_elem, replacements)
    body = doc.element.body
    sect_pr = body.find(qn("w:sectPr"))
    if sect_pr is None:
        body.append(new_elem)
    else:
        body.insert(list(body).index(sect_pr), new_elem)
    return new_elem


def clear_body(doc):
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def add_p(doc, text="", style="Normal"):
    p = doc.add_paragraph(style=style)
    if text:
        p.add_run(text)
    return p


def add_caption(doc, text):
    return add_p(doc, text, "Normal")


def add_model_table(doc, style_name=None):
    rows = [
        ("模型名称", "说明"),
        ("XGBoost", "基于人工提取充电特征的树模型回归基线，用于验证静态特征对SOH的基础表征能力。"),
        ("LSTM", "仅利用循环序列建模的时序网络基线，用于验证跨循环状态依赖的建模效果。"),
        ("CNN-LSTM", "单尺度卷积与LSTM结合的深度学习基线，用于对比多尺度特征提取的增益。"),
        ("MS-CNN-LSTM", "引入多尺度卷积分支但不加入物理一致性约束的结构消融模型。"),
        ("PI-MS-CNNLSTM", "在多尺度特征提取和LSTM时序建模基础上加入软单调物理一致性约束的本文方法。"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    if style_name:
        try:
            table.style = style_name
        except Exception:
            pass
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = val
    return table


def add_paragraphs(doc, texts):
    for text in texts:
        add_p(doc, text, "Normal")


def post_process_text(doc):
    replacements = {
        "自建CRT数据集": "自建数据集",
        "自建CRT退化数据集": "自建数据集",
        "CRT数据集": "自建数据集",
        "CRT电池": "自建数据集电池",
        "自建数据集数据集": "自建数据集",
        "4.2.3 数据预处理与划分": "数据预处理与划分",
        "图4-3  不完全生命周期监督条件下训练电池的生命周期覆盖示意图": "图4-1  部分生命周期监督条件下训练电池的生命周期覆盖示意图",
        "图4-3展示了自建数据集中": "图4-1展示了自建数据集中",
        "图4-4  PI-CNNLSTM整体框架示意图": "图4-2  多尺度物理一致性SOH估计方法整体框架示意图",
        "图4-5  锂离子电池在生命周期早期的容量回升现象示意图": "图4-3  锂离子电池在生命周期早期的容量回升现象示意图",
        "图4-5展示": "图4-3展示",
        "图4-1  HUST数据集中77节锂离子电池的容量衰减轨迹": "图4-4  HUST数据集中77节锂离子电池的容量衰减轨迹",
        "图4-1用于展示": "图4-4用于展示",
        "图4-2  HUST数据集中多源特征与容量退化之间的相关性矩阵": "图4-5  HUST数据集中多源特征与容量退化之间的相关性矩阵",
        "表4-5 在完整监督比例下SOH的估算RMSE（%）对比": "表4-5 架构与物理约束二维消融矩阵下的RMSE（%）对比",
        "图4-4  自建数据集容量退化轨迹与标定节点分布": "图4-8  自建数据集容量退化轨迹与标定节点分布",
        "图4-4进一步展示了自建数据集中": "图4-8进一步展示了自建数据集中",
        "图4-12 轨迹偏差异常分数及异常检测结果": "图4-12 轨迹偏差异常分数及异常风险提示结果",
        "异常检测结果": "异常风险提示结果",
        "异常检测性能": "异常风险提示性能",
        "异常检测方案": "异常风险提示方案",
        "异常检测": "异常风险提示",
    }

    exact_text_replacements = {
        "图4-4  PI-CNNLSTM整体框架示意图": "图4-2  多尺度物理一致性SOH估计方法整体框架示意图",
        "图4-8  不同监督比例下各模型MAE变化趋势对比": "图4-9  不同监督比例下各模型MAE变化趋势对比",
        "图4-9  部分监督条件下未见电池上的SOH预测轨迹对比": "图4-10  部分监督条件下未见电池上的SOH预测轨迹对比",
        "图4-10  部分监督条件下未见电池上的预测误差演化对比": "图4-11  部分监督条件下未见电池上的预测误差演化对比",
        "图4-11 正常与故障SOH预测轨迹对比": "图4-12 正常与故障SOH预测轨迹对比",
        "图4-12 轨迹偏差异常分数及异常风险提示结果": "图4-13 轨迹偏差异常分数及异常风险提示结果",
        "图4-11展示了代表性电池自建数据集中某代表性电池": "图4-12展示了自建数据集中某代表性电池",
        "表4-4  完整监督条件下（）模型预测性能对比": "表4-4  完整监督条件下（r=1.0）模型预测性能对比",
    }

    for p in doc.paragraphs:
        original = p.text
        text = original
        for old, new in replacements.items():
            text = text.replace(old, new)
        for old, new in exact_text_replacements.items():
            text = text.replace(old, new)
        if text == original:
            continue
        has_drawing = bool(p._p.findall(".//" + qn("w:drawing")))
        has_pict = bool(p._p.findall(".//" + qn("w:pict")))
        has_omath = bool(p._p.findall(".//" + qn("m:oMath")))
        if has_drawing or has_pict or has_omath:
            for run in p.runs:
                if run.text:
                    run.text = ""
            for run in p.runs:
                if run.text == "":
                    run.text = text
                    break
        else:
            p.text = text


def main():
    orig = Document(str(SRC))
    template = Document(str(SRC))
    clear_body(template)

    children = body_children(orig)
    table_style = None
    try:
        table_style = orig.tables[2].style.name
    except Exception:
        table_style = None

    repl = {
        "自建CRT退化数据集": "面向实际运行场景采集的自建电池退化数据集",
        "自建CRT数据集": "自建数据集",
        "CRT数据集": "自建数据集",
        "CRT-B03": "自建数据集中某代表性电池",
        "4块CRT电池": "4块自建数据集电池",
        "基于异常电池数据的SOH轨迹异常识别扩展": "工程扩展：基于预测轨迹的异常风险提示",
        "异常检测方案": "异常风险提示方案",
        "异常识别": "异常风险提示",
    }

    # Title
    add_p(template, "第四章 基于多尺度特征与物理一致性约束的锂离子电池SOH估计方法", "Heading 1")

    # 4.1
    add_p(template, "4.1 面向实际BMS的SOH估计问题背景", "Heading 1")
    add_paragraphs(
        template,
        [
            "锂离子电池健康状态（State of Health, SOH）的准确估计是电池管理系统（Battery Management System, BMS）进行安全预警、寿命管理和运维决策的重要基础。在实际运行中，BMS能够持续记录电压、电流、温度、时间以及充放电阶段信息，这些运行数据具有连续性强、采集成本低和部署便利等特点，是在线SOH估计最直接的数据来源。",
            "与运行数据相比，真实容量和SOH标签并不能在完整生命周期内连续获得。容量标定通常依赖周期性容量测试、离线诊断或特定维护窗口，不仅耗时，而且会受到工况安排、测试成本和设备条件的限制。因此，实际工程中更常见的数据形态并不是“每个循环均有SOH标签”，而是“运行特征连续可得，健康状态标签稀疏可得”。",
            "这一数据现实决定了SOH估计模型不能只面向理想完整监督条件设计。若模型训练完全依赖密集SOH标签，则在真实BMS部署中会受到标定频率不足、标签区间不连续以及电池个体差异的影响，导致模型在无标注生命周期区间的预测稳定性下降。由此，如何在保留完整运行轨迹的同时利用有限SOH标签进行稳健建模，成为本章关注的核心问题。",
            "从公开验证角度看，HUST锂离子电池数据集具有统一充电协议、多阶段放电工况和完整生命周期容量记录，能够较好反映实际运行中负载变化对退化轨迹的影响。本文首先选取HUST公开数据集作为基线验证平台，用于检验所提出方法在公开基准条件下的基础估计能力、跨电池泛化能力以及多尺度结构的有效性。",
            "从工程应用角度看，仅在公开数据集上验证仍不足以说明模型对实际标定受限场景的适用性。为进一步贴近BMS中容量标定稀疏、不连续的部署条件，本文结合面向实际运行场景采集的自建电池退化数据集，构建部分生命周期监督验证场景，用于分析模型在有限SOH标签条件下的稳定性与物理一致性。",
            "因此，本章采用“工程需求—问题建模—方法设计—公开验证—工程验证”的组织路线。首先从实际BMS数据现实出发定义部分生命周期监督问题；随后提出基于多尺度特征提取与物理一致性约束的SOH估计方法；再分别通过HUST公开数据集和自建数据集完成基础性能验证与工程标定受限场景验证；最后探索利用SOH预测轨迹进行异常风险提示的工程扩展。",
        ],
    )

    # 4.2
    add_p(template, "4.2 部分生命周期监督问题建模", "Heading 1")
    add_paragraphs(
        template,
        [
            "面向实际BMS部署条件，本文将SOH估计任务定义为部分生命周期监督问题。该问题的关键并不是简单减少训练样本，而是在完整保留电池运行特征轨迹的前提下，仅对缺失SOH标注的循环屏蔽数据误差项，使模型能够同时利用有限标签和连续运行过程信息。",
            "设第i个电池的生命周期长度为Ti，其运行特征序列由各循环或充电片段中提取的多维特征构成。完整监督条件下，每个循环均具有对应SOH标签；而在部分生命周期监督条件下，仅有部分循环集合具有可用SOH标定结果，其余循环虽然保留电压、电流、温度、时间等运行特征，但不提供SOH数值监督。",
        ],
    )
    # Keep original formula-related paragraphs from the problem modeling part.
    for idx in [27, 28, 29, 30, 31, 32, 33]:
        append_clone(template, children[idx], repl)
    add_paragraphs(
        template,
        [
            "上述Label Masking机制的意义在于，训练过程并不删除无标签循环，而是仅使其不参与数据误差计算。这样可以避免退化序列被人为截断，使模型仍能观察同一电池在完整生命周期内的运行特征变化，并为后续物理一致性约束在无标签区间发挥结构性弱监督作用提供条件。",
        ],
    )
    # Label masking figure
    for idx in [34, 35, 36]:
        append_clone(
            template,
            children[idx],
            {
                **repl,
                "图4-3  不完全生命周期监督条件下训练电池的生命周期覆盖示意图": "图4-1  部分生命周期监督条件下训练电池的生命周期覆盖示意图",
                "图4-3展示": "图4-1展示",
            },
        )

    # 4.3
    add_p(template, "4.3 多尺度物理一致性SOH估计方法", "Heading 1")
    add_paragraphs(
        template,
        [
            "针对部分生命周期监督条件下标签稀疏、退化阶段不完整和无标签区间缺少直接监督的问题，本文提出一种基于多尺度特征提取与物理一致性约束的SOH估计方法。该方法的设计不是单纯增加网络复杂度，而是围绕实际BMS数据现实中的两个矛盾展开：一方面，有限SOH标签要求模型具备更强的退化特征表征能力；另一方面，无标签生命周期区间需要借助物理先验获得结构性训练信号。",
        ],
    )
    for idx in range(51, 81):
        append_clone(
            template,
            children[idx],
            {
                **repl,
                "4.3.1 整体框架概述": "4.3.1 整体框架概述",
                "MS-PI-CNNLSTM多尺度物理一致性神经网络架构设计": "多尺度物理一致性SOH估计方法",
                "图4-4  PI-CNNLSTM整体框架示意图": "图4-2  多尺度物理一致性SOH估计方法整体框架示意图",
                "图4-5  锂离子电池在生命周期早期的容量回升现象示意图": "图4-3  锂离子电池在生命周期早期的容量回升现象示意图",
                "图4-5展示": "图4-3展示",
                "4.2.2节": "4.2节",
            },
        )

    # 4.4
    add_p(template, "4.4 公开数据集基线验证", "Heading 1")
    add_paragraphs(
        template,
        [
            "为避免仅在自建数据集上验证导致结论缺乏公开可比性，本文首先选取HUST公开锂离子电池数据集作为基线验证平台。该数据集包含完整生命周期容量记录、统一充电协议和多阶段放电工况，能够用于检验模型在公开基准条件下的基础SOH估计能力、跨电池泛化能力以及多尺度结构的有效性。",
        ],
    )
    add_p(template, "4.4.1 HUST数据集与实验设置", "Heading 2")
    for idx in list(range(13, 23)) + list(range(41, 49)):
        append_clone(
            template,
            children[idx],
            {
                **repl,
                "图4-1": "图4-4",
                "图4-2": "图4-5",
                "表4-1": "表4-1",
                "表4-2": "表4-2",
                "自建部分生命周期": "部分生命周期",
            },
        )
    add_p(template, "4.4.2 对比模型与训练配置", "Heading 2")
    for idx in [84, 85, 86, 87]:
        append_clone(template, children[idx], repl)
    add_model_table(template, table_style)
    for idx in [89]:
        append_clone(template, children[idx], repl)
    add_p(template, "4.4.3 完整监督基准性能", "Heading 2")
    for idx in range(91, 102):
        append_clone(template, children[idx], repl)
    add_p(template, "4.4.4 多尺度结构与物理约束消融", "Heading 2")
    for idx in range(103, 114):
        append_clone(
            template,
            children[idx],
            {
                **repl,
                "表4-5 在完整监督比例下SOH的估算RMSE（%）对比": "表4-5 架构与物理约束二维消融矩阵下的RMSE（%）对比",
                "2×2交叉对比实验": "二维交叉对比实验",
            },
        )

    # 4.5
    add_p(template, "4.5 面向工程数据的部分监督验证与分析", "Heading 1")
    add_paragraphs(
        template,
        [
            "在完成公开数据集基线验证后，本文进一步结合面向实际运行场景采集的自建数据集，构建更接近实际BMS标定条件的部分生命周期监督场景。与完整监督不同，该场景假设运行特征可以连续获得，而SOH标签仅在部分生命周期区间可用，用于检验模型在工程标定受限条件下是否仍能维持稳定的SOH估计能力。",
            "自建数据集以锂离子电池单体为研究对象，通过循环充放电实验记录电池退化过程中的电压、电流、温度、时间及容量变化信息。实验过程中仅在部分循环节点进行容量标定，以获得对应循环的SOH标签。该数据结构更接近实际BMS中“运行特征连续可观测、健康状态间歇标定”的应用场景。",
        ],
    )
    for idx in [37, 38, 39, 40]:
        append_clone(
            template,
            children[idx],
            {
                **repl,
                "图4-4": "图4-8",
                "CRT": "自建数据集",
            },
        )
    add_p(template, "4.5.1 部分监督鲁棒性评估", "Heading 2")
    for idx in range(115, 137):
        append_clone(
            template,
            children[idx],
            {
                **repl,
                "图4-8": "图4-9",
                "图4-9": "图4-10",
                "图4-10": "图4-11",
                "表4-7": "表4-7",
                "表4-8": "表4-8",
            },
        )

    # 4.6
    add_p(template, "4.6 工程扩展：基于预测轨迹的异常风险提示", "Heading 1")
    add_paragraphs(
        template,
        [
            "在实际应用中，SOH估计结果不仅可用于健康状态评估，也可为异常退化风险提示提供辅助信息。本节不试图构建完整故障诊断系统，而是在不额外训练故障诊断模型的前提下，利用已训练SOH估计模型输出的预测轨迹构造偏差信号，用于对突发老化、容量拐点和析锂台阶突降等具有明显轨迹扰动特征的场景形成低成本风险提示。",
        ],
    )
    for idx in range(138, 154):
        append_clone(
            template,
            children[idx],
            {
                **repl,
                "图4-11": "图4-12",
                "图4-12": "图4-13",
                "表4-9": "表4-9",
                "异常检测结果": "异常风险提示结果",
                "异常检测性能": "异常风险提示性能",
                "完整故障诊断模型": "专门故障诊断模型",
            },
        )

    # 4.7
    add_p(template, "4.7 本章小结", "Heading 1")
    add_paragraphs(
        template,
        [
            "本章围绕实际电池管理系统中运行数据连续可采集、SOH标定稀疏且不连续的工程矛盾，提出并验证了一种基于多尺度特征与物理一致性约束的锂离子电池SOH估计方法。",
            "在问题建模方面，本章从BMS数据现实出发，将实际工程中常见的“运行特征连续可得、健康状态标签间歇可得”场景抽象为部分生命周期监督问题。通过Label Masking机制，训练阶段保留完整运行特征轨迹，仅对缺失SOH标注的循环屏蔽数据误差项，从而避免无标签循环被删除后造成退化序列断裂。",
            "在方法设计方面，本文方法由多尺度卷积特征提取、循环时序建模和软单调物理一致性约束组成。多尺度卷积模块用于同时提取局部响应、阶段趋势和长期退化信息；LSTM用于建模跨循环状态演化；软单调约束则在SOH标签缺失区间提供退化方向先验，使无标签生命周期片段也能参与结构性弱监督。",
            "在公开验证方面，本章首先基于HUST公开数据集检验模型的基础估计能力与跨电池泛化能力。实验结果表明，多尺度特征提取能够提升模型对不同退化阶段和不同时间粒度特征的表达能力，物理一致性约束则有助于抑制预测轨迹中的非物理波动。",
            "在工程验证方面，本章进一步结合自建数据集构建部分生命周期监督场景，分析不同SOH标定比例下模型性能的变化规律。结果表明，当监督比例降低时，物理一致性约束的价值更加突出，其作用不在于无条件提升全监督点估计精度，而在于为标签稀缺区间提供结构性弱监督，提升预测轨迹的稳定性和物理合理性。",
            "此外，本章将SOH预测轨迹进一步用于异常退化风险提示。基于预测轨迹偏差的扩展实验表明，该方案在突发老化和析锂台阶突降等轨迹扰动明显的场景中具有一定工程参考价值，但对渐进型异常仍存在检测边界。因此，该部分更适合作为低成本风险提示工具，而非专门故障诊断模型的替代方案。",
            "总体来看，本章的核心并非单纯追求模型结构复杂化，而是围绕实际BMS数据条件构建从问题建模、方法设计到公开验证和工程验证的完整路线，为标定受限条件下的锂离子电池SOH稳健估计提供了一种可解释、可验证的技术方案。",
        ],
    )

    # References, keep original tail.
    add_p(template, "参考文献", "Heading 1")
    for idx in range(165, 206):
        append_clone(template, children[idx], repl)

    post_process_text(template)
    template.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
