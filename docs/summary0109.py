from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_meeting_report():
    doc = Document()
    
    # --- Title ---
    title = doc.add_heading('组会文献汇报：面向碎片化充电数据的锂电池 SOH 估计前沿进展', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # --- Metadata ---
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('汇报人：畅畅 | 日期：2026年1月')
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(100, 100, 100)
    
    doc.add_paragraph("-------------------------------------------------------------------------------------------")

    # --- Section 1: Background ---
    doc.add_heading('一、 综述背景与核心痛点 (Motivation)', level=1)
    
    p = doc.add_paragraph()
    p.add_run('现实背景：').bold = True
    p.add_run(' 传统的 SOH 估计依赖完整的充电曲线（CC-CV），但在实车云端应用中，数据往往是碎片化的。')
    
    p = doc.add_paragraph()
    p.add_run('核心挑战：').bold = True
    doc.add_paragraph('1. 时间轴对齐难：不知道碎片数据对应完整曲线的哪个时间点。', style='List Bullet')
    doc.add_paragraph('2. 物理特征丢失：关键的健康特征（如 ICA 峰值）因数据截断而未被记录。', style='List Bullet')
    doc.add_paragraph('3. 行为随机性：用户里程焦虑导致起始/结束电量高度随机。', style='List Bullet')

    # --- Section 2: Literature Review ---
    doc.add_heading('二、 重点文献详细汇报 (Key Literature)', level=1)

    # Paper 1
    doc.add_heading('1. 基于用户行为心理学的场景分类 (Tang et al., IEEE TEC, 2024)', level=2)
    p = doc.add_paragraph()
    p.add_run('创新点：').bold = True
    p.add_run(' 将数据的随机缺失归因于“用户行为心理学”。')
    doc.add_paragraph('场景定义：根据“里程焦虑”和“停驻时间”定义典型场景（如 Case IV 掐头去尾，Case IX 高压补电）。', style='List Bullet')
    doc.add_paragraph('数据处理：使用 Lorentz-Arctan 公式拟合碎片电压，提取电化学参数（峰面积、位置）作为特征。', style='List Bullet')
    doc.add_paragraph('借鉴意义：为我们的“随机丢弃实验”提供了坚实的理论依据（即模拟真实用户行为）。', style='List Bullet')

    # Paper 2
    doc.add_heading('2. 基于相空间变换的 SOH 估计 (Wang et al., Energy, 2025)', level=2)
    p = doc.add_paragraph()
    p.add_run('创新点：').bold = True
    p.add_run(' 解决碎片数据“时间不对齐”的难题。')
    doc.add_paragraph('数据处理：摒弃 V-t（电压-时间）曲线，构建 v-Δv（电压-电压变化量）坐标系。', style='List Bullet')
    doc.add_paragraph('核心优势：具有“时间不变性”，无论从何时开始充电，特征形状固定。', style='List Bullet')
    doc.add_paragraph('借鉴意义：可作为我们模型的“第二视域”输入，解决时间序列截断问题。', style='List Bullet')

    # Paper 3
    doc.add_heading('3. 基于部分充电曲线“重构”的方法 (Sun et al., IEEE TPEL, 2025)', level=2)
    p = doc.add_paragraph()
    p.add_run('创新点：').bold = True
    p.add_run(' “低压重构高压”的数据修复思想。')
    doc.add_paragraph('方法论：利用 CNN 将低相关性的低压片段（Low-voltage）映射为高相关性的高压片段（High-voltage）。', style='List Bullet')
    doc.add_paragraph('借鉴意义：证明了在物理特征缺失时，利用深度学习进行“跨域重构”的可行性。', style='List Bullet')

    # Paper 4
    doc.add_heading('4. 基于显式数学拟合的特征提取 (Lai et al., Batteries, 2024)', level=2)
    p = doc.add_paragraph()
    p.add_run('创新点：').bold = True
    p.add_run(' 轻量级、可解释的全局特征提取。')
    doc.add_paragraph('数据处理：提出 Log-Linear 模型 (A*ln(t) + Bt + C) 拟合电压曲线。', style='List Bullet')
    doc.add_paragraph('借鉴意义：即使数据破碎，拟合出的系数 (A, B, C) 仍能反映全局极化特性，适合作为 PINN 的物理约束对象。', style='List Bullet')
    
    # Paper 5
    doc.add_heading('5. 基于退化模式识别的分组预测 (Fu et al., Batteries, 2025)', level=2)
    p = doc.add_paragraph()
    p.add_run('创新点：').bold = True
    p.add_run(' “先分类，后预测”的分治思想。')
    doc.add_paragraph('方法论：利用聚类算法识别电池不同的衰退模式，训练专属模型。', style='List Bullet')
    doc.add_paragraph('借鉴意义：佐证了我们要进行的“基于用户行为分类”建模的合理性。', style='List Bullet')

    # --- Section 3: Project Connection ---
    doc.add_heading('三、 总结与本项目创新点融合 (Proposed Method)', level=1)
    
    p = doc.add_paragraph('基于上述调研，本项目 (CNN-LSTM-PINN) 的改进方向如下：')
    
    doc.add_paragraph('场景定义视域：引用 Tang et al. (2024)，将“随机丢弃”实验包装为“基于用户里程焦虑的稀疏采样测试”。', style='List Number')
    doc.add_paragraph('双视域特征融合：', style='List Number')
    sub_p1 = doc.add_paragraph('视域 A (显式物理)：利用 Lai/Tang 的数学拟合提取电化学参数。', style='List Bullet 2')
    sub_p1.paragraph_format.left_indent = Pt(36)
    sub_p2 = doc.add_paragraph('视域 B (隐式动力学)：利用 Wang 的 v-Δv 变换提取相空间特征。', style='List Bullet 2')
    sub_p2.paragraph_format.left_indent = Pt(36)
    doc.add_paragraph('PINN 核心价值：在数据缺失/稀疏采样的“空白区”，传统 LSTM 会发生预测漂移，而 PINN 利用物理方程约束退化轨迹的单调性和平滑性，实现精准外推。', style='List Number')

    # Save
    file_path = "/mnt/data/Group_Meeting_Report_SOH_Estimation.docx"
    doc.save(file_path)
    return file_path

# Generate the document
file_path = create_meeting_report()
file_path