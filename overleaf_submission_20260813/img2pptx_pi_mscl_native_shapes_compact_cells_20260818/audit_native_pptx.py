from __future__ import annotations

import io
import json
import zipfile
from collections import Counter
from pathlib import Path

import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(__file__).resolve().parent
SOURCE_DIR = ROOT.parent / "img2pptx_pi_mscl_proportional_cells_20260818"
PPTX_PATH = ROOT / "PI-MSCL_architecture_native_shapes_compact_cells_20260818.pptx"
PPTX_UNGROUPED = ROOT / "PI-MSCL_architecture_native_shapes_fully_ungrouped_compact_cells_20260818.pptx"
REFERENCE = SOURCE_DIR / "qa" / "rendered_svg.png"
PREVIEW = ROOT / "qa" / "pptx_slide_preview.png"
PREVIEW_UNGROUPED = ROOT / "qa" / "pptx_slide_preview_fully_ungrouped.png"
QA_DIR = ROOT / "qa"


def walk_shapes(shapes, parent: str = "slide"):
    for shape in shapes:
        record = {
            "name": shape.name,
            "parent": parent,
            "shape_type": str(shape.shape_type),
            "is_group": shape.shape_type == MSO_SHAPE_TYPE.GROUP,
            "has_text_frame": bool(getattr(shape, "has_text_frame", False)),
            "text": shape.text if getattr(shape, "has_text_frame", False) else "",
        }
        yield record
        if record["is_group"]:
            yield from walk_shapes(shape.shapes, shape.name)


def image_audit():
    reference = Image.open(REFERENCE).convert("RGB")
    preview = Image.open(PREVIEW).convert("RGB")
    if preview.size != reference.size:
        preview = preview.resize(reference.size, Image.Resampling.LANCZOS)
    ref = np.asarray(reference, dtype=np.int16)
    got = np.asarray(preview, dtype=np.int16)
    delta = np.abs(ref - got)
    mae = float(delta.mean())
    rmse = float(np.sqrt(np.mean((ref - got) ** 2)))
    changed_gt_12 = float(np.mean(np.max(delta, axis=2) > 12))
    amplified = np.clip(delta * 4, 0, 255).astype(np.uint8)
    Image.fromarray(amplified).save(QA_DIR / "native_vs_svg_amplified_diff.png")
    Image.blend(reference, preview, 0.5).save(QA_DIR / "native_vs_svg_overlay.png")
    reference.save(QA_DIR / "reference_svg_render.png")
    ungrouped = Image.open(PREVIEW_UNGROUPED).convert("RGB")
    if ungrouped.size != preview.size:
        ungrouped = ungrouped.resize(preview.size, Image.Resampling.LANCZOS)
    grouped_vs_ungrouped = np.abs(
        np.asarray(preview, dtype=np.int16) - np.asarray(ungrouped, dtype=np.int16)
    )
    return {
        "reference": str(REFERENCE),
        "preview": str(PREVIEW),
        "mae_rgb": round(mae, 4),
        "rmse_rgb": round(rmse, 4),
        "fraction_pixels_max_channel_diff_gt_12": round(changed_gt_12, 6),
        "grouped_vs_fully_ungrouped_mae_rgb": round(float(grouped_vs_ungrouped.mean()), 6),
    }


def package_audit(pptx_path: Path):
    with zipfile.ZipFile(pptx_path) as zf:
        slide_xml = zf.read("ppt/slides/slide1.xml").decode("utf-8")
        media = [n for n in zf.namelist() if n.startswith("ppt/media/")]
    return {
        "p_sp": slide_xml.count("<p:sp>"),
        "p_cxnSp": slide_xml.count("<p:cxnSp>"),
        "p_grpSp": slide_xml.count("<p:grpSp>"),
        "p_pic": slide_xml.count("<p:pic>"),
        "a_t": slide_xml.count("<a:t>"),
        "a_custGeom": slide_xml.count("<a:custGeom>"),
        "media_files": media,
    }


def main():
    QA_DIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation(PPTX_PATH)
    inventory = list(walk_shapes(prs.slides[0].shapes))
    type_counts = Counter(item["shape_type"] for item in inventory)
    names = [item["name"] for item in inventory]
    text_items = [item for item in inventory if item["has_text_frame"] and item["text"]]

    # Prove text frames are writable without altering the delivered file.
    edit_probe = False
    if text_items:
        probe_prs = Presentation(PPTX_PATH)
        probe_shapes = list(walk_shapes(probe_prs.slides[0].shapes))
        target_name = text_items[0]["name"]

        def set_same_text(shapes):
            for shape in shapes:
                if shape.name == target_name and getattr(shape, "has_text_frame", False):
                    old = shape.text
                    shape.text = old
                    return True
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP and set_same_text(shape.shapes):
                    return True
            return False

        if set_same_text(probe_prs.slides[0].shapes):
            buffer = io.BytesIO()
            probe_prs.save(buffer)
            edit_probe = len(buffer.getvalue()) > 0

    package = package_audit(PPTX_PATH)
    ungrouped_package = package_audit(PPTX_UNGROUPED)
    visual = image_audit()
    report = {
        "pptx": str(PPTX_PATH),
        "slides": len(prs.slides),
        "top_level_objects": len(prs.slides[0].shapes),
        "recursive_object_count": len(inventory),
        "recursive_type_counts": dict(type_counts),
        "named_objects": len(names),
        "unique_names": len(set(names)),
        "editable_text_objects": len(text_items),
        "in_memory_text_edit_probe": edit_probe,
        "package": package,
        "fully_ungrouped_package": ungrouped_package,
        "visual": visual,
        "hard_checks": {
            "single_slide": len(prs.slides) == 1,
            "no_picture_objects": package["p_pic"] == 0,
            "no_embedded_media": package["media_files"] == [],
            "at_least_500_native_objects": package["p_sp"] + package["p_cxnSp"] >= 500,
            "all_svg_texts_are_editable_text_frames": len(text_items) >= 87 and package["a_t"] >= 87,
            "all_objects_have_unique_names": len(names) == len(set(names)),
            "semantic_module_groups_present": package["p_grpSp"] >= 20,
            "text_edit_probe_succeeded": edit_probe,
            "visual_mae_below_8": visual["mae_rgb"] < 8.0,
            "grouped_and_ungrouped_render_identically": visual["grouped_vs_fully_ungrouped_mae_rgb"] < 0.05,
            "fully_ungrouped_has_zero_groups": ungrouped_package["p_grpSp"] == 0,
            "fully_ungrouped_has_no_pictures": ungrouped_package["p_pic"] == 0,
        },
    }
    report["all_hard_checks_pass"] = all(report["hard_checks"].values())
    (QA_DIR / "native_editability_inventory.json").write_text(
        json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (QA_DIR / "native_pptx_final_audit.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if not report["all_hard_checks_pass"]:
        raise SystemExit(2)


if __name__ == "__main__":
    main()
